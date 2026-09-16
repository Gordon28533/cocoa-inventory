"""
Two further corrections.

1. Table count. The schema in Backend/lib/database.js creates exactly five
   tables: departments, users, inventory, requisitions, audit_logs. A
   multi-item requisition is several rows in `requisitions` sharing a batch_id;
   there is no requisition_items table. The dissertation says "seven core
   tables".

2. "Tamper-evident". audit_logs is an ordinary table. The application only
   INSERTs, so it is append-only by convention, but there is no trigger, no
   revocation of UPDATE/DELETE, and no hash chain. "Tamper-evident" claims a
   property the schema does not provide.
"""
from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
doc = Document(DOC)
ps = doc.paragraphs


def set_text(p, text):
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def replace_in_runs(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    joined = "".join(r.text for r in p.runs)
    if old in joined:
        set_text(p, joined.replace(old, new))
        return True
    return False


# ---------------------------------------------------- report: what Appendix A says
print("=== Appendix A table mentions ===")
start = next(i for i, p in enumerate(ps) if p.text.strip().startswith("APPENDIX A"))
end = next(i for i, p in enumerate(ps) if p.text.strip().startswith("APPENDIX B"))
for i in range(start, end):
    t = ps[i].text.strip()
    for name in ("departments", "users", "inventory", "requisitions",
                 "audit_logs", "requisition_items"):
        if f"CREATE TABLE" in t and name in t:
            print(f"  [{i}] CREATE TABLE … {name}")
            break
counts = [i for i in range(start, end)
          if "core table" in ps[i].text or "seven" in ps[i].text.lower()]
for i in counts:
    print(f"  [count claim {i}] {ps[i].text.strip()[:140]}")

# ---------------------------------------------------- report: para 44
print("\n=== paragraph 44 (before) ===")
print(ps[44].text[:700])

# ---------------------------------------------------- fixes
print("\n=== applying fixes ===")

ok = replace_in_runs(
    ps[427],
    "The normalized database schema, documented across seven core tables,",
    "The normalized database schema, documented across five core tables,")
print(f"  [427] 'seven core tables' -> 'five core tables' : {ok}")

ok = replace_in_runs(
    ps[427],
    "requisition lifecycle management, and tamper-evident audit logging",
    "requisition lifecycle management, and append-only audit logging")
print(f"  [427] 'tamper-evident' -> 'append-only'         : {ok}")

ok = replace_in_runs(
    ps[44],
    "a retrievable, tamper-evident record",
    "a retrievable, append-only record")
print(f"  [44]  'tamper-evident' -> 'append-only'         : {ok}")

doc.save(DOC)

# ---------------------------------------------------- verify
d2 = Document(DOC)
txt = "\n".join(p.text for p in d2.paragraphs)
print("\n=== residual check ===")
for term in ["tamper-evident", "seven core tables", "requisition_items"]:
    print(f"  {'STILL PRESENT' if term in txt else 'clear        '}  {term}")
print("\n=== paragraph 44 (after) ===")
print(d2.paragraphs[44].text[:420])
