from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
doc = Document(DOC)
ps = doc.paragraphs


def set_text(p, text):
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def replace_in_runs(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    joined = "".join(r.text for r in p.runs)
    if old in joined:
        set_text(p, joined.replace(old, new))
        return True
    return False


print("=== [239] entity model: describe what was actually built ===")
set_text(ps[239],
    "The entity relationship model underlying the schema is presented in Figure "
    "3.4. The implemented schema comprises five tables — departments, users, "
    "inventory, requisitions and audit_logs — connected by one-to-many "
    "relationships. The many-to-many relationship between requisitions and "
    "inventory items is not resolved by a separate junction table. Instead, a "
    "multi-item request is stored as one row per line item in the requisitions "
    "table, with all rows of a single submission sharing a batch_id. This was a "
    "deliberate departure from the conventional junction-table design: because "
    "each line item carries its own status and its own approver columns, a "
    "requisition can in principle be approved or rejected item by item, which a "
    "shared header row would not permit. The trade-off is that request-level "
    "attributes such as the department and the originating flags are repeated "
    "on every line of a batch.")
print("  rewritten")

print("\n=== [44] audit log: remove the unsupported traceability claim ===")
ok = replace_in_runs(ps[44],
    "The append-only audit log table structure means that historical records "
    "cannot be modified or deleted without a clearly traceable administrative "
    "action.",
    "The audit log is written only by insertion — the application never updates "
    "or deletes an entry — so the record of actions accumulates as a complete "
    "history. Restricting modification at the database role level, or chaining "
    "entries cryptographically, would be required before alteration by a "
    "privileged user could be detected; this is identified as a recommendation "
    "in Chapter Five.")
print(f"  replaced: {ok}")

doc.save(DOC)

d2 = Document(DOC)
txt = "\n".join(p.text for p in d2.paragraphs)
print("\n=== residual check ===")
for term in ["requisition_items", "eight core entities", "tamper-evident",
             "seven core tables",
             "cannot be modified or deleted without a clearly traceable"]:
    print(f"  {'STILL PRESENT' if term in txt else 'clear        '}  {term}")
