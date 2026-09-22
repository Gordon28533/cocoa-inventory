from docx import Document
DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
d = Document(DOC)
ps = d.paragraphs

KEYS = ["Workflow", "workflow path", "Accounts Manager", "Deputy HOD",
        "IT Manager", "approval path", "Path 1", "Path 2", "Path 3",
        "branch_account", "ho_account", "it_approved", "hod_approved"]

seen = set()
for i, p in enumerate(ps):
    t = p.text
    if any(k in t for k in KEYS) and len(t.strip()) > 40 and i not in seen:
        seen.add(i)
        print(f"\n--- idx {i} [{p.style.name}] ---")
        print(t.strip()[:1000])

print("\n\n=== Word tables mentioning approval roles ===")
for ti, tbl in enumerate(d.tables):
    txt = " ".join(c.text for r in tbl.rows for c in r.cells)
    if "Accounts Manager" in txt or "IT Manager" in txt or "Deputy" in txt:
        print(f"\nTABLE {ti} ({len(tbl.rows)} rows x {len(tbl.columns)} cols)")
        for r in tbl.rows:
            print("   | " + " | ".join(c.text.strip()[:38] for c in r.cells))
