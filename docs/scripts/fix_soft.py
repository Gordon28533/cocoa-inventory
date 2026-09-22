from docx import Document
DOC = r'C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx'
doc = Document(DOC)

def set_text(p, text):
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)

new = ('The Inventory Service owns the authoritative source of truth for stock '
       'levels, holding a single quantity per inventory item; there is no '
       'separate reserved quantity, because stock is deducted at fulfilment '
       'rather than held back at submission. It exposes RESTful endpoints for '
       'browsing, searching, creating, updating, and deleting inventory items. '
       'The fulfilment deduction is performed inside a transaction, with the '
       'contended row locked for the duration, so concurrent fulfilments of the '
       'same item cannot interleave.')

ok = False
for p in doc.paragraphs:
    if 'soft-reserved quantity for each inventory item' in p.text:
        set_text(p, new); ok = True; break
print('patched:', ok)
doc.save(DOC)

d2 = Document(DOC)
txt = '\n'.join(p.text for p in d2.paragraphs)
for t in ['soft-reserved','soft reserved','reserved quantity for each']:
    print(('STILL PRESENT  ' if t in txt else 'clear          ') + t)
