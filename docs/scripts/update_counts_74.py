"""Four branch-guard regression tests were added: backend is now 41, total 74."""
from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
doc = Document(DOC)
ps = doc.paragraphs


def replace_in_runs(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    joined = "".join(r.text for r in p.runs)
    if old in joined:
        for r in list(p.runs)[1:]:
            r._element.getparent().remove(r._element)
        if p.runs:
            p.runs[0].text = joined.replace(old, new)
        else:
            p.add_run(joined.replace(old, new))
        return True
    return False


EDITS = [
    ("All seventy cases pass — 37 backend and 33 frontend.",
     "All seventy-four cases pass — 41 backend and 33 frontend."),
    ("Both suites pass in full: 37 of 37 backend cases and 33 of 33 frontend cases, "
     "70 in total.",
     "Both suites pass in full: 41 of 41 backend cases and 33 of 33 frontend cases, "
     "74 in total."),
    ("C.1 Backend suite — 37 of 37 passing",
     "C.1 Backend suite — 41 of 41 passing"),
    ("Seventy automated test cases pass — 37 backend and 33 frontend —",
     "Seventy-four automated test cases pass — 41 backend and 33 frontend —"),
]

print("=== updating counts ===")
for old, new in EDITS:
    done = False
    for p in ps:
        if replace_in_runs(p, old, new):
            print(f"  ok   {old[:58]}…")
            done = True
            break
    if not done:
        print(f"  MISS {old[:58]}…")

doc.save(DOC)

d2 = Document(DOC)
txt = "\n".join(p.text for p in d2.paragraphs)
print("\n=== residual stale figures (all should be clear) ===")
for term in ["37 backend", "37 of 37", "seventy cases", "70 in total",
             "Seventy automated", "sixty-six", "47 functional"]:
    print(f"  {'STILL PRESENT' if term in txt else 'clear        '}  {term}")

print("\n=== current test sentences ===")
for p in d2.paragraphs:
    if "41" in p.text and "backend" in p.text.lower():
        print("  •", p.text.strip()[:190])
