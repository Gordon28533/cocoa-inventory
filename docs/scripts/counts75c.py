import re
from docx import Document
DOC = r'C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx'
doc = Document(DOC)
n=0
for p in doc.paragraphs:
    j = ''.join(r.text for r in p.runs)
    if 'Seventy-four' in j or 'seventy-four' in j:
        k = j.replace('Seventy-four','Seventy-five').replace('seventy-four','seventy-five')
        for r in list(p.runs)[1:]: r._element.getparent().remove(r._element)
        if p.runs: p.runs[0].text = k
        else: p.add_run(k)
        n+=1
doc.save(DOC); print('paragraphs patched:', n)

d2=Document(DOC); txt='\n'.join(p.text for p in d2.paragraphs)
for t in ['eventy-four','orty-one','41 backend','74 in total','41 of 41']:
    print(('STILL  ' if t in txt else 'clear  ')+t)
print()
for p in d2.paragraphs:
    if re.search(r'(eventy-five|42 of 42|42 backend)', p.text):
        print(' *', p.text.strip()[:200]); print()
