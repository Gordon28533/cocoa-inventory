"""
Round 3: the remaining table descriptions in Section 3.7.1.

  Claim                                   Reality (Backend/lib/database.js)
  --------------------------------------  ----------------------------------------
  Users: "email address"                  No email column exists.
  Users: "timestamps for account          Only created_at. No updated_at.
   creation and last modification"
  Users: role "constrained ... through a  role is VARCHAR(50). Permitted values
   database-level ENUM constraint"        live in Backend/lib/roles.js, enforced
                                          in application code.
  Items: "soft-reserved quantity" and     inventory has a single quantity column.
   "computed available quantity"          No reservation model exists.
  Items: "the department or location       No department or location column.
   responsible for each item"
  Audit: "no UPDATE or DELETE             No GRANT or REVOKE statement exists
   permissions granted to the              anywhere in the codebase. Verified by
   application database user, enforcing    grep across Backend/.
   its immutability at the database
   authorization level"
"""
from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
doc = Document(DOC)


def set_text(p, text):
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def rewrite(marker, new):
    for p in doc.paragraphs:
        if marker in p.text:
            set_text(p, new)
            return True
    return False


EDITS = [
("The role field is constrained to the eight permitted role values through a database-level ENUM constraint",
 "The Users table stores each staff member's surrogate identifier, full name, "
 "Staff ID, bcrypt-hashed password, department foreign key, assigned role, an "
 "active flag, and the account creation timestamp. Both the full name and the "
 "Staff ID carry uniqueness constraints, so no two accounts can share either. "
 "The role field is a VARCHAR; its eight permitted values are defined in "
 "Backend/lib/roles.js and enforced in application code rather than by a "
 "database ENUM type or CHECK constraint, which means role validity depends on "
 "every write path applying that check."),

("soft-reserved quantity, reorder threshold, unit of measure",
 "The Items table stores inventory item identifiers, names, descriptions, "
 "category and type classifications, the quantity currently held, the unit of "
 "measure, the reorder threshold at which replenishment is indicated, and "
 "timestamps for creation and last update. Stock is represented by a single "
 "quantity column. The design holds no separate reserved quantity, and "
 "therefore no derived available quantity, because stock is deducted at "
 "fulfilment rather than reserved at submission — a requisition may be raised "
 "for more than is in stock, and the shortfall is detected when Stores attempts "
 "to fulfil it. A CHECK constraint prevents the quantity from falling below "
 "zero."),

("The table has no UPDATE or DELETE permissions granted to the application database user",
 "The Audit Logs table is an append-only ledger recording significant system "
 "events. Each record captures the event timestamp, the authenticated user "
 "identifier, the action performed, the requisition affected where one applies, "
 "and a free-text details column carrying contextual metadata such as a "
 "rejection reason. \"Append-only\" describes how the application uses the "
 "table — every write is an INSERT, and no code path updates or deletes a row — "
 "rather than a restriction enforced by the database: UPDATE and DELETE have "
 "not been revoked for the application's database role, so a holder of those "
 "credentials could alter the log without detection. Revoking those privileges, "
 "or chaining each row to a hash of its predecessor, is the work required to "
 "make the log tamper-evident rather than merely append-only, and is identified "
 "as a recommendation in Chapter Five."),
]

print("=== round 3 corrections ===")
for marker, new in EDITS:
    print(f"  {'ok  ' if rewrite(marker, new) else 'MISS'}  {marker[:64]}…")

doc.save(DOC)

d2 = Document(DOC)
txt = "\n".join(p.text for p in d2.paragraphs)
print("\n=== residual check ===")
for term in ["email address", "soft-reserved", "available quantity, derived",
             "ENUM constraint", "no UPDATE or DELETE permissions",
             "database authorization level", "last modification"]:
    print(f"  {'STILL PRESENT' if term in txt else 'clear        '}  {term}")
print(f"\nparagraphs: {len(d2.paragraphs)}   images: {len(d2.inline_shapes)}")
