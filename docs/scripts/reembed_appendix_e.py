"""
Replace the ten Appendix E figures with freshly captured screenshots.

The chrome change altered every page's height, so the new captures have
different aspect ratios from the ones embedded. Swapping the image bytes alone
would stretch them, so each figure's extent is recomputed: desktop figures keep
their 6.00in width, mobile figures keep their 7.20in height, and any figure
that would then exceed the usable page height is scaled down to fit.
"""
import pathlib, struct
from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
SHOTS = pathlib.Path(r"C:\Users\PC\cocoa-inventory\docs\screenshots")
EMU = 914400
NS_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

# Figure order in the document -> capture file
FILES = [
    "01-login.png", "02-inventory-list.png", "03-stock-alerts.png",
    "04-requisition-form.png", "05-my-requisitions.png", "06-departments.png",
    "07-user-management.png", "08-audit-log.png",
    "09-mobile-dashboard.png", "10-mobile-login.png",
]
MOBILE = {"09-mobile-dashboard.png", "10-mobile-login.png"}

DESKTOP_W_IN = 6.00
MOBILE_H_IN = 7.20
MAX_H_IN = 8.00   # leaves room for the caption beneath


def png_size(path):
    """Width/height from the PNG IHDR, without needing an imaging library."""
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"not a PNG: {path}")
    w, h = struct.unpack(">II", head[16:24])
    return w, h


doc = Document(DOC)
ps = doc.paragraphs
start = next(i for i, p in enumerate(ps) if p.text.strip().startswith("APPENDIX E"))
end = next((i for i, p in enumerate(ps) if i > start and p.style.name == "Heading 1"), len(ps))

inlines = []
for i in range(start, end):
    for inl in ps[i]._p.findall(f".//{NS_WP}inline"):
        inlines.append((i, inl))

print(f"found {len(inlines)} figures in Appendix E; {len(FILES)} captures to place")
assert len(inlines) == len(FILES), "figure count does not match capture count"

sect = doc.sections[0]
text_h_in = (sect.page_height - sect.top_margin - sect.bottom_margin) / EMU
text_w_in = (sect.page_width - sect.left_margin - sect.right_margin) / EMU

seen_rids = {}
for (para_idx, inl), fname in zip(inlines, FILES):
    path = SHOTS / fname
    px_w, px_h = png_size(path)
    aspect = px_w / px_h

    if fname in MOBILE:
        h_in = MOBILE_H_IN
        w_in = h_in * aspect
    else:
        w_in = DESKTOP_W_IN
        h_in = w_in / aspect

    if h_in > MAX_H_IN:                       # too tall: constrain by height
        h_in = MAX_H_IN
        w_in = h_in * aspect
    if w_in > text_w_in:                      # too wide: constrain by width
        w_in = text_w_in
        h_in = w_in / aspect

    cx, cy = int(w_in * EMU), int(h_in * EMU)

    # Swap the image bytes on the related part.
    blip = inl.find(f".//{NS_A}blip")
    rid = blip.get(f"{NS_R}embed")
    if rid in seen_rids:
        raise RuntimeError(f"rId {rid} shared by two figures ({seen_rids[rid]} and {fname})")
    seen_rids[rid] = fname
    part = doc.part.related_parts[rid]
    part._blob = path.read_bytes()

    # Resize: both the inline extent and the picture's own extent.
    ext = inl.find(f"{NS_WP}extent")
    ext.set("cx", str(cx)); ext.set("cy", str(cy))
    for a_ext in inl.findall(f".//{NS_A}ext"):
        a_ext.set("cx", str(cx)); a_ext.set("cy", str(cy))

    print(f"  para {para_idx:<5} {fname:<26} {px_w}x{px_h}px -> "
          f"{w_in:.2f} x {h_in:.2f} in")

doc.save(DOC)
print("\nsaved.")

# ---------------------------------------------------------------- verify
d2 = Document(DOC)
print(f"images in document: {len(d2.inline_shapes)}  (expected 17)")
over = []
for sh in d2.inline_shapes:
    w, h = sh.width / EMU, sh.height / EMU
    if w > text_w_in + 0.01 or h > text_h_in + 0.01:
        over.append((round(w, 2), round(h, 2)))
print(f"figures exceeding the {text_w_in:.2f} x {text_h_in:.2f} in text area: "
      f"{over if over else 'none'}")
