"""Final cross-check: dissertation vs deck vs Q&A vs code."""
import re, pathlib
from docx import Document
from pptx import Presentation

REPO = pathlib.Path(r"C:\Users\PC\cocoa-inventory")
DISS = Document(REPO / "CHAPTER ONE-FIVE (CORRECTED).docx")
QA   = Document(REPO / "Defence Q&A Preparation.docx")
DECK = Presentation(REPO / "deck" / "CMC Inventory System - Defence (GCTU Template).pptx")

diss = "\n".join(p.text for p in DISS.paragraphs)
qa   = "\n".join(p.text for p in QA.paragraphs)
deck = "\n".join(sh.text_frame.text for s in DECK.slides for sh in s.shapes
                 if sh.has_text_frame)
code = "\n".join(p.read_text(encoding="utf-8", errors="ignore")
                 for p in (REPO / "Backend").rglob("*.js"))

print("=" * 72)
print("A.  FABRICATION SWEEP — none of these should appear anywhere")
print("=" * 72)
BANNED = ["requisition_items", "version stamping", "OCC with", "47 functional",
          "94% line", "100% coverage", "tamper-evident", "soft-reserved",
          "Notification Service", "parent department references",
          "seven core tables", "eight core entities", "Rostami",
          "Kumar Anugula", "two outstanding failures", "reserved_qty"]
for term in BANNED:
    hits = [n for n, t in (("dissertation", diss), ("deck", deck), ("Q&A", qa))
            if term in t]
    print(f"  {'FOUND IN ' + ','.join(hits) if hits else 'clear':<28} {term}")

print()
print("=" * 72)
print("B.  TEST COUNTS — must agree everywhere")
print("=" * 72)
for name, txt in (("dissertation", diss), ("deck", deck), ("Q&A", qa)):
    nums = sorted(set(re.findall(r"\b(\d{2})\b(?=[^.]{0,40}(?:backend|frontend|test))",
                                 txt, re.I)))
    words = [w for w in ["seventy-four", "Seventy-four", "seventy", "sixty-six"]
             if w in txt]
    print(f"  {name:<14} two-digit near 'test': {nums}   words: {words}")

print()
print("=" * 72)
print("C.  WORKFLOW CHAINS — deck vs dissertation vs code")
print("=" * 72)
code_transitions = re.findall(
    r'first\.status === "(\w+)".*?user\.role === "(\w+)"', code)
print("  transitions found in requisitionRoutes.js:")
for st, role in dict.fromkeys(code_transitions):
    print(f"     status={st:<24} role={role}")
print()
print("  dissertation Appendix B summary:")
for p in DISS.paragraphs:
    if "The resulting paths are" in p.text:
        print("     " + p.text.strip()[:330])
print()
print("  deck slide 10:")
for s in DECK.slides:
    if s.shapes.title is not None and "APPROVAL WORKFLOWS" in s.shapes.title.text_frame.text:
        for sh in s.shapes:
            if sh.has_text_frame and "→" in sh.text_frame.text:
                for line in sh.text_frame.text.split("\n"):
                    if "→" in line:
                        print("     " + line.strip())

print()
print("=" * 72)
print("D.  SCHEMA FACTS — claim vs code")
print("=" * 72)
checks = [
    ("5 tables created",      len(re.findall(r"CREATE TABLE IF NOT EXISTS", code)) == 5),
    ("no requisition_items",  "requisition_items" not in code),
    ("CHECK quantity >= 0",   "quantity >= 0" in code),
    ("staffId UNIQUE",        re.search(r'"staffId".*UNIQUE', code) is not None),
    ("status is VARCHAR",     re.search(r"status\s+VARCHAR", code) is not None),
    ("no ENUM type",          "CREATE TYPE" not in code),
    ("no GRANT/REVOKE",       not re.search(r"\b(GRANT|REVOKE)\b", code)),
    ("no email column",       not re.search(r"email\s+VARCHAR", code)),
    ("8 roles in roles.js",   len(re.findall(r'"\w+",?\s*$',
                              (REPO / "Backend/lib/roles.js").read_text(), re.M)) >= 8),
    ("role re-read per req",  "loadLiveUser" in code),
    ("FOR UPDATE lock",       "FOR UPDATE" in code),
    ("isHO guard on approve", re.search(r'user\.role === "deputy_hod"\)\s*&&', code)
                              is not None or "isHO\n      ) {" in code),
]
for label, ok in checks:
    print(f"  {'PASS' if ok else 'FAIL'}   {label}")

print()
print("=" * 72)
print("E.  DOCUMENT INTEGRITY")
print("=" * 72)
print(f"  dissertation paragraphs : {len(DISS.paragraphs)}")
print(f"  dissertation images     : {len(DISS.inline_shapes)}")
print(f"  deck slides             : {len(DECK.slides)}")
heads = [p.text.strip() for p in DISS.paragraphs if p.style.name == "Heading 1"]
print(f"  Heading 1 order (tail)  : {heads[-7:]}")
refs = sum(1 for p in DISS.paragraphs
           if re.match(r"^[A-Z][A-Za-z'’\-]+,\s+[A-Z]\.|^(The PostgreSQL|OWASP|Ghana Cocoa|van der)",
                       p.text.strip()))
print(f"  reference entries       : {refs}")
