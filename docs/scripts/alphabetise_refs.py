"""APA orders the list alphabetically by the first author's surname, ignoring a
leading 'The'. OWASP (O) currently sits after PostgreSQL (P). Reorder in place."""
from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
doc = Document(DOC)
ps = doc.paragraphs

STARTS = (
    "Banks, A.", "Casciaro, M.", "Chopra, S.", "Dumas, M.", "Elmasri, R.",
    "Ghana Cocoa Board.", "Kleppmann, M.", "Laudon, K.", "Newman, S.",
    "The PostgreSQL Global Development Group.", "OWASP Foundation.",
    "Pressman, R.", "Richards, G.", "Silberschatz, A.", "Sommerville, I.",
    "Stallings, W.", "van der Aalst, W.", "Wild, T.",
)

heading = [p for p in ps if p.text.strip() == "REFERENCES"
           and p.style.name.startswith("Heading")][0]
entries = [p for p in ps if p.text.strip().startswith(STARTS)]


def sort_key(p):
    t = p.text.strip()
    if t.startswith("The "):
        t = t[4:]
    return t.lower()


print("before:")
for p in entries:
    print("   ", p.text.strip()[:58])

ordered = sorted(entries, key=sort_key)

anchor = heading._p
for p in reversed(ordered):
    anchor.addnext(p._p)

doc.save(DOC)

d2 = Document(DOC)
h = [i for i, p in enumerate(d2.paragraphs) if p.text.strip() == "REFERENCES"
     and p.style.name.startswith("Heading")][0]
print("\nafter:")
for p in d2.paragraphs[h + 1:h + 20]:
    if p.text.strip():
        print("   ", p.text.strip()[:58])
