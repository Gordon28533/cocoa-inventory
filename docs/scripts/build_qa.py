"""Build the defence Q&A preparation document."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
OUT = r"C:\Users\PC\cocoa-inventory\Defence Q&A Preparation.docx"

# ---- first, report on the tamper-evident claim ------------------------------
d0 = Document(SRC)
hits = [(i, p.text) for i, p in enumerate(d0.paragraphs) if "tamper" in p.text.lower()]
print(f"'tamper' occurrences in the dissertation: {len(hits)}")
for i, t in hits:
    idx = t.lower().find("tamper")
    print(f"  [{i}] …{t[max(0,idx-150):idx+150]}…")

# ---- build the Q&A doc ------------------------------------------------------
doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(0.9)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)

RED = RGBColor(0xA6, 0x1B, 0x1B)
GREY = RGBColor(0x55, 0x55, 0x55)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def q(text):
    p = doc.add_paragraph()
    p.space_before = Pt(10)
    p.space_after = Pt(2)
    r = p.add_run("Q.  " + text)
    r.bold = True
    return p


def a(text, danger=False):
    p = doc.add_paragraph()
    p.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.28)
    r = p.add_run(text)
    if danger:
        r.font.color.rgb = RED
    return p


def note(text):
    p = doc.add_paragraph()
    p.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.28)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    return p


# ---------------------------------------------------------------- title
t = doc.add_heading("Project Defence — Anticipated Questions", level=0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("A Web-Based Inventory and Requisition Management System for "
              "Cocoa Marketing Company (Ghana) Limited\nGodwin Ahiamadia  ·  "
              "Prepared 16 September 2026")
r.font.size = Pt(10.5)
r.font.color.rgb = GREY

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run(
    "This is ordered by risk, not by chapter. Section 1 covers the questions "
    "that could actually expose a weakness in the project — answer these badly "
    "and it costs you marks. Sections 2 and 3 are the routine questions. "
    "Section 4 is the material you want them to ask about. Section 5 is what "
    "to do before you walk in.")
r.font.size = Pt(10.5)

# ================================================================ SECTION 1
h("1.  The questions that can hurt you", 1)

note("Each of these targets something genuinely soft in the project. The "
     "answers given are honest ones — do not claim more than these say.")

q("Can a Head of Department approve a branch requisition? Walk me through what "
  "happens.")
a("FIXED — AND THIS IS NOW ONE OF YOUR BEST ANSWERS.")
a("Suggested answer: \"No, and that's enforced. The transition from pending to "
  "hod_approved tests three things: the requisition's status, the approver's "
  "role, and whether the requisition originated at head office. An HOD has no "
  "step in the branch chain, so the attempt is refused with a 403 that names "
  "the chain which does apply — Branch: Accounts, then Accounts Manager.\"")
a("Then volunteer how you found it: \"It was actually a defect until I checked "
  "the implementation against my own workflow design. The guard originally "
  "tested role and status but not origin, so an HOD could pull a branch "
  "requisition into the head-office chain — and because Branch Accounts only "
  "accepts requisitions at status pending, and nothing ever returns one to "
  "pending, that approval step was bypassed permanently. It was invisible "
  "because every approval test in the suite used a head-office requisition. I "
  "added four regression tests: two proving the branch case is refused, and two "
  "proving the fix didn't over-correct and block the legitimate approvers.\"")
note("This is a strong answer because it shows you can find a defect that the "
     "tests, the code review and the running system all failed to surface. If "
     "they ask how you verified the fix: the two refusal tests were confirmed to "
     "fail with the fix reverted, so they genuinely test the guard rather than "
     "passing by accident.")

q("Log in as someone in the IT department and raise a requisition. Which "
  "approval chain does it take?")
a("THE BRANCH CHAIN — AND THAT IS ALMOST CERTAINLY NOT WHAT YOU INTEND.", danger=True)
a("The seed data creates eight departments and flags only 'Head Office' as "
  "is_head_office = 1. IT, HR, Finance and Stores are all seeded with 0, so the "
  "system treats them as branches. A requisition from the IT department is "
  "therefore routed Accounts then Accounts Manager — it never reaches the HOD, "
  "and an IT item from the IT department gets no IT Manager review.")
a("If your demo raises a requisition from any of those four departments, the "
  "behaviour will contradict your workflow slide in front of the examiner.")
a("Suggested answer if asked: \"Those four are functional departments that sit "
  "at head office, so they should carry the head-office flag; the seed sets it "
  "only on the department literally named 'Head Office'. It's a seed-data "
  "classification error rather than a fault in the routing logic — the routing "
  "reads the flag correctly, the flag is set wrongly for those rows.\"")
note("Decide this before you demo: either demonstrate from a Takeover Center "
     "department (genuinely a branch, behaves correctly) or from Head Office, "
     "and avoid IT/HR/Finance/Stores. The durable fix is to set is_head_office "
     "= 1 for those four rows — but only you know CMC's real structure, so "
     "confirm it before changing it.")

q("A branch office orders a laptop. Does the IT Manager review it?")
a("NO — AND THIS IS A REAL GAP IN THE ROUTING RULES.", danger=True)
a("The three chains are selected by two flags. The branch chain is chosen on "
  "is_head_office alone and does not consult is_it_item, so a branch IT "
  "requisition goes Accounts then Accounts Manager and never reaches the IT "
  "Manager. Only head-office IT items get the technical review.")
a("Suggested answer: \"No, and it should. The routing tests the head-office flag "
  "first, so a branch requisition takes the branch path whatever the item is — "
  "which means an IT item ordered by a branch gets financial approval but no "
  "technical review. It's a gap in the rules rather than a deliberate "
  "exemption. I've recorded it as a limitation, and the fix is to let the "
  "is_it_item flag insert the IT Manager step irrespective of origin.\"")
note("Your document now records this as a limitation, so you are consistent if "
     "they check. Naming it before they do is much stronger than conceding it "
     "afterwards.")

q("What makes your audit log tamper-evident?")
a("BE CAREFUL — THE DOCUMENT OVERSTATES THIS.", danger=True)
a("audit_logs is an ordinary table. The application only ever INSERTs into it, "
  "so it is append-only by convention, but there is no trigger, no database "
  "rule revoking UPDATE and DELETE, and no hash chain. Anyone with database "
  "credentials could alter a row and nothing would detect it.")
a("Suggested answer: \"It's append-only in the sense that the application never "
  "updates or deletes entries — every action writes a new row with the actor, "
  "the action, the requisition and a timestamp. It is not cryptographically "
  "tamper-evident: someone with direct database access could alter it "
  "undetected. Making it genuinely tamper-evident would mean either revoking "
  "UPDATE and DELETE on that table at the database role level, or chaining each "
  "row to a hash of the previous one. That's a recommendation rather than "
  "something I implemented.\"")

q("You claim the system is safe under concurrent access. How do you know?")
a("Suggested answer: \"Three mechanisms, and one honest limitation. The "
  "availability read takes a row-level lock with SELECT ... FOR UPDATE, so a "
  "second fulfilment of the same item waits rather than reading stock that is "
  "about to be taken. The decrement then repeats the availability condition in "
  "its own WHERE clause and uses RETURNING to confirm a row actually matched, "
  "so a deduction can never rest on a stale read. And a CHECK (quantity >= 0) "
  "constraint makes a negative quantity unrepresentable regardless of "
  "application logic. The limitation is that I verified this with a regression "
  "test that simulates the interleaving — it is a logical test of the guard, "
  "not a load test. I have not measured throughput under sustained contention, "
  "and Chapter Four says so.\"")

q("Has CMC actually used this system? What did the staff say?")
a("Do not overstate this. If you have not run it with real staff on real data, "
  "say so plainly.")
a("Suggested answer: \"No. The system is deployed and fully functional, and I "
  "exercised every workflow end to end with test data, but it has not been run "
  "against live CMC operational data and there has been no formal user "
  "acceptance testing with staff. That is the first limitation on my "
  "limitations slide, and my first recommendation is a pilot with a single "
  "department before any wider rollout.\"")

q("Where do you store the authentication token, and what are the implications?")
a("Suggested answer: \"In the browser's localStorage. That is a deliberate "
  "trade-off rather than an oversight: it survives a page refresh and keeps the "
  "API stateless, but it is readable by any JavaScript running on the page, so "
  "a cross-site scripting flaw would expose the token. The more secure option "
  "is an httpOnly cookie, which JavaScript cannot read, at the cost of needing "
  "CSRF protection. For a production deployment I would move to httpOnly "
  "cookies.\"")
note("The dissertation previously claimed the opposite — that the token was "
     "held in memory to avoid this exact risk. That was corrected. Do not "
     "revert to the old claim under pressure.")

q("Show me your Business Rules Engine.")
a("Suggested answer: \"It is the routing logic that runs at submission: it "
  "reads the department's head-office flag and the IT-item flag on each "
  "requisition and selects one of the three approval chains from them. I use "
  "the term to describe that function — it is routing logic within the "
  "requisition module, not a separately configurable rules engine of the kind "
  "you would find in a commercial BPM product.\"")
note("If your supervisor expects a configurable engine, describe it as 'the "
     "routing logic' throughout rather than defending the grander term.")

q("Your literature review discusses Petri nets and workflow soundness. Did you "
  "verify your workflow is sound?")
a("Suggested answer: \"No, and the document now says so explicitly. Petri net "
  "formalism informed how I thought about the state machine — that every "
  "requisition should reach a terminal state and that there should be no dead "
  "transitions — but I did not carry out formal soundness verification. That "
  "was outside the scope of the project.\"")

# ================================================================ SECTION 2
h("2.  Section-by-section questions", 1)

h("Background and problem statement", 2)
for text, ans in [
    ("Why did you choose this topic?",
     "Personal exposure to the problem, and CMC is a real organisation with a "
     "real multi-site stores function whose approval rules differ by site and "
     "by item type — which makes it more than a CRUD exercise."),
    ("Of the five deficiencies you listed, which is the most serious and why?",
     "Unenforced approval authority and unchecked issue of stock — those two "
     "carry financial exposure. The others cost time rather than money."),
    ("How did you establish that these problems actually exist?",
     "Be precise about what you did: observation of the paper procedure, "
     "informal discussion, document review. If you did not run structured "
     "interviews or a survey, say so — inventing a methodology you did not "
     "follow is the single easiest way to get caught."),
    ("What is the cost of the current system to CMC?",
     "If you do not have figures, do not invent them. \"I did not quantify it "
     "in cedis; I characterised it qualitatively as delay, inaccuracy and lack "
     "of accountability.\""),
]:
    q(text); a(ans)

h("Aim and objectives", 2)
for text, ans in [
    ("Are your objectives measurable?",
     "Five of the six are demonstrable by inspection — the schema exists, the "
     "roles exist, the workflows enforce, the concurrency guard holds, the "
     "system is deployed. Objective one, analysing the existing procedure, is "
     "evidenced by Chapter Three."),
    ("Take me through each objective and tell me whether you met it.",
     "Rehearse this. It is the most commonly asked question in a defence and "
     "the answer is Section 5.2 of your Chapter Five. Six objectives, each with "
     "one sentence of evidence."),
    ("Which objective was hardest?",
     "Objective five — preventing over-issue under concurrency. It is the only "
     "part of the system where two users genuinely contend for the same row, "
     "and getting it right needed a lock, a guarded update and a constraint "
     "rather than any one of them alone."),
]:
    q(text); a(ans)

h("Literature review", 2)
for text, ans in [
    ("What gap did you identify?",
     "Commercial inventory modules assume a single approval chain. None of "
     "those reviewed model a chain that switches on both item type and "
     "originating site, which is exactly CMC's requirement."),
    ("Which existing systems did you review, and how do they compare?",
     "Odoo, BarCloud and the Cayuse routing platform. Know one concrete "
     "strength and one concrete gap for each."),
    ("Why is RBAC appropriate here rather than access control lists?",
     "Permissions map onto organisational roles that already exist — HOD, "
     "Stores, Accounts — so roles are the natural unit. Per-user ACLs would "
     "need re-editing every time a person changed post."),
]:
    q(text); a(ans)

h("Methodology", 2)
for text, ans in [
    ("Why Agile rather than Waterfall?",
     "The approval rules were not fully known at the outset — they emerged by "
     "working through real cases. Waterfall assumes you can specify completely "
     "before building, which was not true here. That is the honest answer and "
     "it is the correct one."),
    ("How many iterations, and what was in each?",
     "Four: authentication and roles; inventory and alerts; requisitions and "
     "approval chains; fulfilment, stock deduction and audit logging."),
    ("Did you work alone? How did you apply Agile as an individual?",
     "Be straightforward — iterative increments with working software at the "
     "end of each, not ceremonies like stand-ups that require a team."),
]:
    q(text); a(ans)

h("Systems analysis and design", 2)
for text, ans in [
    ("Explain your architecture and justify three tiers.",
     "React SPA on Vercel, Express REST API on Render, PostgreSQL on Neon. "
     "Separation lets the presentation layer be replaced or a mobile client "
     "added without touching business logic, and keeps the database "
     "unreachable from the browser."),
    ("Why PostgreSQL rather than MySQL?",
     "A strong answer — you actually migrated. Stricter type checking, better "
     "constraint support including the CHECK constraint the concurrency design "
     "depends on, and RETURNING, which lets the guarded decrement confirm in "
     "one statement that a row matched."),
    ("Walk me through your ER diagram. What normal form is your schema in?",
     "Five tables: departments, users, inventory, requisitions, audit_logs. "
     "Foreign keys run from users and requisitions to departments, and from "
     "audit_logs to users. Note what you should NOT claim: requisitions.item_id "
     "has no foreign key to inventory, so referential integrity between a "
     "requisition line and its item is not enforced by the database. If asked "
     "why, the honest answer is that it was not added; adding it would be a "
     "one-line improvement."),
    ("Where is the junction table between requisitions and inventory items?",
     "There isn't one, and that is deliberate. A multi-item request is stored "
     "as one row per line item sharing a batch_id, so each line carries its own "
     "status and approver columns and could be decided individually — which a "
     "header row plus junction table would not allow. The cost is that "
     "request-level attributes repeat on every line. Be aware that Figure 3.4 "
     "still shows a junction table; if they point at it, say the figure is the "
     "original conceptual model and the implemented design is the one described "
     "in the text."),
    ("Why does a requisition store a department name as well as a department "
     "id? Isn't that redundant?",
     "A fair challenge — it is denormalised. Justify it as a historical "
     "snapshot so a renamed department does not rewrite past records, or "
     "concede it if you cannot."),
    ("Why three workflows rather than one configurable chain?",
     "Three fixed chains match the three rules the organisation actually "
     "operates, and an explicit state machine is verifiable by reading it. A "
     "fully configurable engine would be more flexible and considerably harder "
     "to prove correct — that is future work."),
    ("What happens if someone tries to approve out of turn?",
     "Refused with 403, and the message names their role, the requisition's "
     "current status and the chain that actually applies — so the refusal tells "
     "them what to do instead of merely saying no."),
]:
    q(text); a(ans)

h("Implementation and testing", 2)
for text, ans in [
    ("How many tests, and what do they cover?",
     "Seventy-five: 42 backend across six files, 33 frontend across eleven. "
     "Authentication, departments, item authorisation, requisitions, users, "
     "server health, and the interface components. Know this number — it is in "
     "the dissertation and on the slide."),
    ("What is your test coverage percentage?",
     "\"I didn't measure it. No coverage instrumentation was run, so I don't "
     "quote a figure. What I can tell you is exactly which behaviours the suite "
     "exercises, and Appendix C also lists the three areas that carry no "
     "automated coverage at all — concurrent access under load, end-to-end "
     "browser testing, and the token refresh paths.\" Never invent a "
     "percentage: they can ask you to run the tool in front of them."),

    ("Do your tables use ENUM types or CHECK constraints for status and role?",
     "Only one CHECK constraint exists, on inventory quantity. Status and role "
     "are VARCHAR columns whose permitted values are enforced in application "
     "code — the approval state machine for status, and roles.js for role. Say "
     "that plainly; it is a reasonable design that you should not dress up as "
     "database-level enforcement. If asked what it costs you: validity then "
     "depends on every write path applying the check."),
    ("Why bcrypt at ten rounds rather than twelve?",
     "Ten balances brute-force resistance against login latency of roughly 80 "
     "to 120 milliseconds. The cost factor is configurable, so it can be raised "
     "as hardware improves without changing code."),
    ("How are passwords stored and how is login handled?",
     "The user is looked up by Staff ID; bcrypt.compare validates the submitted "
     "password against the stored hash. Plaintext is never stored or logged."),
    ("If an administrator demotes a user, when does that take effect?",
     "Immediately. The JWT carries identity, but the role is re-read from the "
     "database on every request, so a token issued before the demotion does not "
     "stay privileged until it expires. This is a genuinely good design "
     "decision — make sure you mention it."),
    ("What happens if two people fulfil the same requisition at once?",
     "The second waits on the row lock, then finds the guarded decrement "
     "matches no row, and the whole transaction rolls back rather than "
     "deducting. Nothing partial is left behind."),
    ("Your tests had drifted and most were failing. Doesn't that mean you "
     "weren't testing during development?",
     "Answer it head on: \"For part of the project, effectively yes. The suite "
     "was not maintained through the authentication change and the PostgreSQL "
     "migration, so it was green on stale assumptions. Repairing it uncovered a "
     "real defect in batch approval that the stale suite had been hiding. The "
     "lesson I take is that an unmaintained test suite is worse than none, "
     "because it produces false confidence.\""),
]:
    q(text); a(ans)

h("Conclusions and recommendations", 2)
for text, ans in [
    ("What are the limitations of your system?",
     "Not run on live CMC data; no offline capability; basic reporting with no "
     "analytics; in-application notifications only. State them without being "
     "prompted — an examiner who has to extract limitations reads them as "
     "concealed."),
    ("What would you do differently if you started again?",
     "Have a real answer ready. Maintaining the test suite continuously, and "
     "settling the approval rules completely before building the state machine, "
     "are both true and both creditable."),
    ("What is the future work?",
     "Pilot with one department; email or SMS notification at each approval "
     "step; a reporting module over the audit log; integration with HR records; "
     "and a security review before production use."),
    ("Could this be deployed at another COCOBOD subsidiary?",
     "Yes, with configuration of departments and roles — the routing rules are "
     "the part that would need revisiting, since they encode CMC's specific "
     "structure."),
]:
    q(text); a(ans)

# ================================================================ SECTION 3
h("3.  Technical questions they may drill into", 1)
for text, ans in [
    ("What is a SQL injection attack and how does your system prevent it?",
     "Every query uses parameterised placeholders rather than string "
     "concatenation, so user input is never interpreted as SQL."),
    ("What does ACID mean, and which property matters most here?",
     "Atomicity — the fulfilment operation deducts stock for every line item, "
     "transitions the status and writes the audit entry inside one transaction. "
     "If any item has insufficient stock, the whole thing rolls back."),
    ("What is a race condition? Give the example from your own system.",
     "Two fulfilments read the same stock level, both see enough, both deduct, "
     "and the total issued exceeds what existed. Check-then-act with a gap "
     "between the check and the act."),
    ("What is the difference between authentication and authorisation?",
     "Authentication establishes who the user is — verifyToken. Authorisation "
     "establishes what they may do — requireRole, plus the workflow guard."),
    ("Why REST rather than GraphQL?",
     "A small, fixed set of resources and operations, where REST's mapping onto "
     "HTTP verbs is simpler to reason about and to secure per endpoint."),
    ("How would this behave with 500 concurrent users?",
     "Do not guess. \"I have not load tested it, so I can't give you a measured "
     "answer. The row lock serialises contention on a single item, which is the "
     "obvious bottleneck, and measuring that is work I would need to do before "
     "production.\""),
    ("What happens if the database connection drops mid-transaction?",
     "The transaction is not committed, so it rolls back — no partial "
     "deduction. The request surfaces an error to the user."),
    ("Why did you migrate from MySQL to PostgreSQL mid-project? What broke?",
     "A good story if you tell it well: type inference on bare parameters, "
     "boolean-to-smallint coercion, and case-sensitive quoted identifiers each "
     "caused real failures that had to be found and fixed."),
]:
    q(text); a(ans)

# ================================================================ SECTION 4
h("4.  Questions you want them to ask", 1)
note("These are where the project is strongest. If you get an open invitation "
     "— \"tell us about an interesting problem you solved\" — go here.")
for text, ans in [
    ("Tell us about a difficult bug you solved.",
     "The batch approval defect. Every test approved using a real numeric "
     "requisition id; the interface only ever sends a placeholder of zero for a "
     "batch. The route validated the id before it read the batch identifier, so "
     "every batch approval failed with \"Invalid requisition ID\" while the "
     "tests stayed green. It is a precise illustration of a suite that "
     "exercises paths the product never takes."),
    ("What are you most proud of in this project?",
     "Choose one and be specific. The concurrency design with its three "
     "independent mechanisms, or the diagnostic 403 that tells an approver "
     "which chain applies rather than simply refusing, are both good answers."),
    ("How did you ensure the system is maintainable?",
     "Separation into routes, library modules and shared helpers; the approval "
     "logic as one explicit readable state machine rather than scattered "
     "conditionals; and an automated test suite that now passes in full."),
]:
    q(text); a(ans)

# ================================================================ SECTION 5
h("5.  Before you walk in", 1)
for item in [
    "Change the admin password. The system is on a public URL and the "
    "default admin123 still works — an examiner who opens the deployed site "
    "could sign in as administrator. This is the one item on this list that "
    "is still outstanding.",
    "Redraw Figure 3.4 if you can. The ER diagram shows a requisition_items "
    "junction table that the implementation does not have; the surrounding "
    "text now describes the real five-table design, but the figure still "
    "shows the original conceptual model.",
    "Fill in the index number and supervisor name on the title slide.",
    "Rehearse the demo once with the network disconnected, so you know what "
    "you would say if it fails. The four screenshots on slide 12 are the "
    "fallback.",
    "Warm up the backend BEFORE you present. This was measured, not "
    "estimated: a cold request to the Render service took 43 seconds. If the "
    "examiner opens your app cold, it will look broken. Load the site "
    "yourself a few minutes beforehand and keep a tab open.",
    "If asked about dependency security: the backend's own runtime packages "
    "are patchable with a non-breaking npm audit fix (express and body-parser "
    "advisories). The critical advisories are in build tooling and in "
    "bcrypt's install-time chain, none of which ship to or run in the "
    "deployed service. Do not run the fix in the days before your defence — "
    "a lockfile change could break a working deployment.",
    "Be able to state the test figure from memory: 75 tests, 42 backend and "
    "33 frontend, across 17 files.",
    "Re-read your own Chapter Five section 5.2. The objective-by-objective "
    "assessment is the most likely single question.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.space_after = Pt(6)
    p.add_run(item)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run(
    "A closing note on tone: every answer above that admits a limitation is "
    "stronger than a claim you cannot support. Examiners are testing whether "
    "you understand your own system, and knowing exactly where it is weak is "
    "the clearest evidence that you do.")
r.italic = True
r.font.size = Pt(10.5)
r.font.color.rgb = GREY

doc.save(OUT)
print("\nsaved:", OUT)
print("paragraphs:", len(doc.paragraphs))
qs = sum(1 for p in doc.paragraphs if p.text.startswith("Q.  "))
print("questions:", qs)
