"""
The 18 reference entries sit at the tail of Appendix D (Glossary), immediately
before the APPENDIX E heading, while the document's own REFERENCES heading near
the end is empty. Move the entries so they follow that heading.

Entries are located by matching their leading author string rather than by a
hardcoded index, so the script is safe to re-run.
"""
from docx import Document

DOC = r"C:\Users\PC\cocoa-inventory\CHAPTER ONE-FIVE (CORRECTED).docx"
doc = Document(DOC)
ps = doc.paragraphs

STARTS = (
    "Banks, A.", "Casciaro, M.", "Chopra, S.", "Dumas, M.", "Elmasri, R.",
    "Ghana Cocoa Board.", "Kleppmann, M.", "Laudon, K.", "Newman, S.",
    "The PostgreSQL Global Development Group.", "OWASP Foundation.",
    "Pressman, R.", "Richards, G.", "Silberschatz, A.", "Sommerville, I.",
    "Stallings, W.", "van der Aalst, W.", "Wild, T.",
)

entries = [p for p in ps if p.text.strip().startswith(STARTS)]
print(f"reference entries found: {len(entries)} (expected 18)")
for p in entries:
    print("   ", p.text.strip()[:74])

heading = None
for p in ps:
    if p.text.strip() == "REFERENCES" and p.style.name.startswith("Heading"):
        heading = p
print(f"\nREFERENCES heading found: {heading is not None}")

pos_before = [i for i, p in enumerate(ps) if p in entries]
h_idx = ps.index(heading)
print(f"entries currently at {pos_before[0]}-{pos_before[-1]}; heading at {h_idx}")

if h_idx < pos_before[0]:
    print("Heading already precedes the entries — nothing to do.")
else:
    anchor = heading._p
    for p in reversed(entries):
        anchor.addnext(p._p)
    doc.save(DOC)
    print("\nmoved and saved.")

# ------------------------------------------------------------------ verify
d2 = Document(DOC)
ps2 = d2.paragraphs
h2 = [i for i, p in enumerate(ps2)
      if p.text.strip() == "REFERENCES" and p.style.name.startswith("Heading")][0]
print(f"\n=== Verification: 12 paragraphs from the REFERENCES heading ({h2}) ===")
for p in ps2[h2:h2 + 12]:
    t = p.text.strip()
    if t:
        print(f"  [{p.style.name}] {t[:80]}")

appE = [i for i, p in enumerate(ps2) if p.text.strip().startswith("APPENDIX E")]
print(f"\nAPPENDIX E heading now at {appE}; REFERENCES at {h2} "
      f"(references should come last: {h2 > appE[0] if appE else 'n/a'})")
print(f"total reference entries under the heading: "
      f"{sum(1 for p in ps2[h2:] if p.text.strip().startswith(STARTS))}")
